# -*- coding: utf-8 -*-
"""Verification renderer for the minimal deck: draws the saved .pptx back to PNGs and reports overflow."""
import os, io
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.normpath(os.path.join(HERE, "..", "Gully_Labs_Strategy_Deck.pptx"))
OUT = os.path.join(HERE, "preview")
os.makedirs(OUT, exist_ok=True)
FR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
FI = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FSER = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"
SCALE = 1.5; PXPI = 96.0
_cache = {}


def font(size, bold, italic, serif):
    key = (round(size * 2), bold, italic, serif)
    if key not in _cache:
        p = (FSER if serif else (FB if bold and not italic else FR))
        if italic and not serif:
            p = FR
        _cache[key] = ImageFont.truetype(p, max(4, int(round(size * PXPI / 72.0 * SCALE))))
    return _cache[key]


def rgb(c, d=(0, 0, 0)):
    try:
        if c.type is not None and c.rgb is not None:
            v = c.rgb; return (v[0], v[1], v[2])
    except Exception:
        pass
    return d


def px(v):
    return Emu(v).inches * PXPI * SCALE


def run_spc(r):
    try:
        rPr = r._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        if rPr is not None and rPr.get('spc'):
            return float(rPr.get('spc')) / 100.0
    except Exception:
        pass
    return 0.0


def is_serif(r):
    return (r.font.name or "") in ("Georgia",)


def wrap(d, runs, width_px):
    chunks = []
    for t, f, col, sp in runs:
        hard = False
        for seg in t.split("\n"):
            for i, w in enumerate(seg.split(" ")):
                chunks.append((("" if i == 0 else " ") + w, f, col, sp, hard))
                hard = False
            hard = True
    lines, cur, curw = [], [], 0
    for w, f, col, sp, hard in chunks:
        ww = d.textlength(w, font=f) + len(w) * sp / 72.0 * PXPI * SCALE
        if cur and (hard or curw + ww > width_px):
            lines.append(cur); cur, curw = [], 0
        cur.append((w, f, col, sp)); curw += ww
    if cur:
        lines.append(cur)
    return lines


def measure(tf, width_in, d):
    total = 0.0
    for p in tf.paragraphs:
        runs = [(r.text, font(r.font.size.pt if r.font.size else 11, bool(r.font.bold), bool(r.font.italic), is_serif(r)),
                 rgb(r.font.color), run_spc(r)) for r in p.runs]
        if not runs:
            total += 11 * 1.2; continue
        ls = p.line_spacing if isinstance(p.line_spacing, (int, float)) else 1.0
        sa = p.space_after.pt if p.space_after is not None else 0
        mx = max((r.font.size.pt if r.font.size else 11) for r in p.runs)
        total += mx * 1.2 * ls * max(1, len(wrap(d, runs, width_in * PXPI * SCALE))) + sa
    return total / PXPI


def main():
    prs = Presentation(PPTX)
    rep = []
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (int(13.3333 * PXPI * SCALE), int(7.5 * PXPI * SCALE)), (244, 239, 230))
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            x, y, w, h = px(sh.left or 0), px(sh.top or 0), px(sh.width or 0), px(sh.height or 0)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pim = Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize((int(w), int(h)))
                    img.paste(pim, (int(x), int(y)))
                except Exception as e:
                    rep.append("S%d picture fail %s" % (idx, e))
                continue
            fill = None; linec = None
            try:
                if sh.fill.type is not None and sh.fill.type == 1:
                    fill = rgb(sh.fill.fore_color)
            except Exception:
                pass
            try:
                if sh.line.fill.type == 1:
                    linec = rgb(sh.line.color)
            except Exception:
                pass
            try:
                oval = "OVAL" in str(sh.auto_shape_type)
            except Exception:
                oval = False
            if fill:
                (d.ellipse if oval else d.rectangle)([x, y, x + w, y + h], fill=fill)
            elif linec:
                (d.ellipse if oval else d.rectangle)([x, y, x + w, y + h], outline=linec, width=max(1, int(SCALE * 0.8)))
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                ml = (tf.margin_left or 0) / 914400.0
                mt = (tf.margin_top or 0) / 914400.0
                inner = max(0.2, w / SCALE / PXPI - 2 * ml)
                need = measure(tf, inner, d)
                if need > h / SCALE / PXPI + 0.05 and "TEXT_BOX" in str(sh.shape_type):
                    rep.append("S%d overflow '%s' need %.2f box %.2f y=%.2f" % (idx, tf.text[:36].replace("\n", " "), need, h / SCALE / PXPI, y / SCALE / PXPI))
                ty = y + mt * PXPI * SCALE
                for p in tf.paragraphs:
                    runs = [(r.text, font(r.font.size.pt if r.font.size else 11, bool(r.font.bold), bool(r.font.italic), is_serif(r)),
                             rgb(r.font.color), run_spc(r)) for r in p.runs]
                    ls = p.line_spacing if isinstance(p.line_spacing, (int, float)) else 1.0
                    sa = (p.space_after.pt if p.space_after is not None else 0)
                    align = str(p.alignment)
                    for ln in wrap(d, runs, inner * PXPI * SCALE):
                        lw = sum(d.textlength(t, font=f) + len(t) * sp / 72.0 * PXPI * SCALE for t, f, c, sp in ln)
                        tx = x + ml * PXPI * SCALE
                        if "CENTER" in align:
                            tx = x + (w - lw) / 2
                        elif "RIGHT" in align:
                            tx = x + w - ml * PXPI * SCALE - lw
                        mf = max(f.size for t, f, c, sp in ln) if ln else 12
                        for t, f, col, sp in ln:
                            d.text((tx, ty), t, font=f, fill=col)
                            tx += d.textlength(t, font=f) + len(t) * sp / 72.0 * PXPI * SCALE
                        ty += mf * 1.22 * ls
                    ty += sa * PXPI * SCALE / 72.0
        img.resize((int(img.width / SCALE), int(img.height / SCALE)), Image.LANCZOS).save(os.path.join(OUT, "slide_%02d.png" % idx))
        print("rendered", idx)
    print("=== issues (%d) ===" % len(rep))
    for r in rep:
        print(" -", r)


if __name__ == "__main__":
    main()
